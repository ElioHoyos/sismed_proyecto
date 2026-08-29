"""Detecta el TIPO de un archivo por sus columnas/hojas, para enrutarlo a su
importador sin que el usuario tenga que saber a cuál va:

    DBF con MES01..MES12                  → ICI (consumo por establecimiento)
    DBF con ALMCOD+MEDLOTE+STKSALDODE     → stock del almacén por lote
    DBF con MEDNOM+MEDTIP+MEDPET          → catálogo de productos (MPRODUCTO)
    XLSX con hoja 'COMPRA …' + Código SISMED → compra centralizada (CENARES)

Solo lee el header (nombres de columnas / hojas), no todos los datos."""
from pathlib import Path

from dbfread import DBF
from openpyxl import load_workbook

TIPO_ICI = "ICI"
TIPO_STOCK_ALMACEN = "STOCK_ALMACEN"
TIPO_CATALOGO = "CATALOGO"
TIPO_CENARES = "CENARES"
TIPO_MOVIM_CAB = "MOVIM_CAB"  # TMOVIM (cabecera de movimientos de kardex)
TIPO_MOVIM_DET = "MOVIM_DET"  # TMOVIMDET (detalle: líneas de producto)
TIPO_DESCONOCIDO = "DESCONOCIDO"


def _campos(ruta: str | Path) -> set[str]:
    tabla = DBF(str(ruta), lowernames=True, ignore_missing_memofile=True)
    return {c.lower() for c in tabla.field_names}


def detectar_tipo(campos: set[str]) -> str:
    # ICI primero: es el único con las 12 columnas de meses (aunque comparte
    # medtip/medpet con el catálogo, no trae MEDNOM).
    if {"mes01", "mes12"} <= campos:
        return TIPO_ICI
    if {"almcod", "medlote", "stksaldode"} <= campos:
        return TIPO_STOCK_ALMACEN
    # Movimientos de kardex (TMOVIM / TMOVIMDET): el detalle trae MEDCOD + ítem;
    # la cabecera trae ALMCODIORG (y MOVREFE, que NO se importa).
    if {"movnumero", "movnumeite", "medcod", "movcantid"} <= campos:
        return TIPO_MOVIM_DET
    if {"movnumero", "almcodiorg", "movcoditip"} <= campos:
        return TIPO_MOVIM_CAB
    if {"mednom", "medtip", "medpet"} <= campos:
        return TIPO_CATALOGO
    return TIPO_DESCONOCIDO


def _detectar_xlsx(ruta: str | Path) -> str:
    """CENARES: alguna hoja se llama 'COMPRA …' y su fila 3 col B es 'Código SISMED'."""
    try:
        wb = load_workbook(ruta, read_only=True)
    except Exception:
        return TIPO_DESCONOCIDO
    try:
        for hoja in wb.sheetnames:
            if "COMPRA" in hoja.upper():
                b3 = wb[hoja].cell(3, 2).value
                if b3 and "SISMED" in str(b3).upper():
                    return TIPO_CENARES
    finally:
        wb.close()
    return TIPO_DESCONOCIDO


def detectar_tipo_archivo(ruta: str | Path) -> str:
    if str(ruta).lower().endswith(".xlsx"):
        return _detectar_xlsx(ruta)
    return detectar_tipo(_campos(ruta))
