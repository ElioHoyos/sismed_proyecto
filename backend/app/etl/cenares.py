"""
Importador del estado de la compra centralizada de CENARES (XLSX).

Hojas 'COMPRA 2025' y 'COMPRA 2026'. Encabezados en la fila 3 (subcabeceras en la
4), datos desde la fila 5. Columnas: B=Código SISMED (llave, zfill 5), C=Código
SIGA, E=Tipo, F=Procedimiento, G/H=Estado situacional (Situación/Observación),
I/J=Registro SIGA (Situación/Observación), K=Contratista, L=Nº contrato,
M=Convocatoria, N=Buena Pro, O=Primera entrega, P=Observación.

La fecha de entrega (O) viene como fecha real o como texto ("NOVIEMBRE 2026"):
se guarda la fecha parseada cuando se puede y SIEMPRE el texto original.

Idempotente: recargar un año reemplaza sus filas (delete + insert), no duplica.
"""
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import delete
from sqlalchemy.orm import Session

from app.core.db import SessionLocal
from app.models.compra_centralizada import CompraCentralizada

ANCHO_COD_SISMED = 5


@dataclass
class ResumenCenares:
    archivo: str
    por_anio: dict[int, int] = field(default_factory=dict)

    def texto(self) -> str:
        detalle = ", ".join(f"{a}: {n}" for a, n in sorted(self.por_anio.items()))
        return f"{self.archivo}: {detalle or 'sin hojas de compra reconocidas'}"


def _texto(v) -> str | None:
    if v is None:
        return None
    t = str(v).strip()
    return t or None


def _fecha(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def _anio_de_hoja(nombre: str) -> int | None:
    m = re.search(r"(20\d\d)", nombre)
    return int(m.group(1)) if m else None


def contar_filas(ruta_xlsx: str | Path) -> int:
    """Cuenta filas con Código SISMED en las hojas de compra — para la
    previsualización, sin importar nada."""
    wb = load_workbook(ruta_xlsx, read_only=True, data_only=True)
    try:
        total = 0
        for hoja in wb.sheetnames:
            if _anio_de_hoja(hoja) is None:
                continue
            for r in wb[hoja].iter_rows(min_row=5, min_col=2, max_col=2, values_only=True):
                if r[0] not in (None, ""):
                    total += 1
        return total
    finally:
        wb.close()


def importar_compra_centralizada(db: Session, ruta_xlsx: str | Path) -> ResumenCenares:
    ruta = Path(ruta_xlsx)
    resumen = ResumenCenares(archivo=ruta.name)
    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        for hoja in wb.sheetnames:
            anio = _anio_de_hoja(hoja)
            if anio is None:
                continue
            ws = wb[hoja]

            # Idempotencia: se reemplaza por completo lo de este año.
            db.execute(delete(CompraCentralizada).where(CompraCentralizada.anio == anio))

            filas: list[CompraCentralizada] = []
            for r in ws.iter_rows(min_row=5, values_only=True):
                def col(i):
                    return r[i] if len(r) > i else None

                cod = _texto(col(1))  # B
                if not cod:
                    continue

                entrega = col(14)  # O
                if isinstance(entrega, (datetime, date)):
                    fecha_ent = _fecha(entrega)
                    texto_ent = fecha_ent.strftime("%d/%m/%Y") if fecha_ent else None
                else:
                    fecha_ent = None
                    texto_ent = _texto(entrega)  # texto tal cual ("NOVIEMBRE 2026")

                filas.append(
                    CompraCentralizada(
                        anio=anio,
                        codigo_sismed=cod.zfill(ANCHO_COD_SISMED),
                        codigo_siga=_texto(col(2)),
                        tipo_producto=_texto(col(4)),
                        procedimiento=_texto(col(5)),
                        estado_situacion=_texto(col(6)),
                        observacion_estado=_texto(col(7)),
                        reg_siga_situacion=_texto(col(8)),
                        reg_siga_observacion=_texto(col(9)),
                        contratista=_texto(col(10)),
                        nro_contrato=_texto(col(11)),
                        fecha_convocatoria=_fecha(col(12)),
                        fecha_buena_pro=_fecha(col(13)),
                        fecha_entrega=fecha_ent,
                        fecha_entrega_texto=texto_ent,
                        observacion=_texto(col(15)),
                    )
                )

            db.add_all(filas)
            resumen.por_anio[anio] = len(filas)

        db.commit()
    finally:
        wb.close()
    return resumen


def main() -> None:
    import sys

    ruta = Path(sys.argv[1])
    db = SessionLocal()
    try:
        r = importar_compra_centralizada(db, ruta)
    finally:
        db.close()
    print(r.texto())


if __name__ == "__main__":
    main()
